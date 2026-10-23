"""A fake OpenAI-compatible server, so no test ever needs a real model."""

import base64
import io
import json
import os
import re
import threading
import time
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

_SZ = re.compile(r"\[SZ(\d+)\]")
_SECTION_INDEX = re.compile(r"=== SECTION section_index (\d+) ===")

SECTION_TITLES = ["Grundbegriffe", "Vorgehen", "Beispiel mit Zahlen", "Grenzen"]
LEAD = "Worum es auf dieser Seite geht."
GIST = ["Die Gerade entsteht aus den Messwerten.", "Der Schwerpunkt trägt sie.",
        "Ein Ausreißer kippt sie."]


def _outline(user: str) -> dict:
    numbers = [int(n) for n in _SZ.findall(user)] or [1]
    groups = [numbers[i::len(SECTION_TITLES)] or [numbers[0]] for i in range(len(SECTION_TITLES))]
    return {
        "title": "Lineare Regression",
        "intro": "Es geht um Geraden. Sie sagen Werte voraus. Der Schwerpunkt trägt sie.",
        "documents": [],
        "sections": [{"section_index": i, "title": title, "description": "Was dazugehört.",
                      "refs": [f"SZ{n}" for n in group[:6]], "start": ""}
                     for i, (title, group) in enumerate(zip(SECTION_TITLES, groups))],
        "lead": LEAD,
        "gist": list(GIST),
        "formula": {"key": "y = a + bx", "text": "Die Gerade."},
        "design": "chapters",
        "hero_prompt": "a calm desk with a ruler",
        "hero_alt": "Ein Schreibtisch mit einem Lineal.",
    }


def _blocks(user: str) -> dict:
    indices = [int(n) for n in _SECTION_INDEX.findall(user)] or [0]
    out = []
    for i in indices:
        out.append({
            "section_index": i,
            "map_label": f"Teil {i + 1}",
            "lead": "Ein Satz zum Abschnitt.",
            "blocks": [
                {"kind": "text",
                 "text": f"Der Abschnitt {i + 1} erklärt seinen Gegenstand in ganzen Sätzen. "
                         "Er nennt die Gründe und die Folgen. Er bleibt bei dem, was im Material steht."},
                {"kind": "bullets", "items": ["Erster Punkt.", "Zweiter Punkt.", "Dritter Punkt."]},
            ],
            "keywords": ["Steigung", "Schwerpunkt", "Residuum"],
        })
    return {"sections": out}


def _topics(user: str) -> dict:
    numbers = [int(n) for n in _SZ.findall(user)] or [1]
    cut = max(1, len(numbers) // 2)
    return {"topics": [
        {"title": "Erstes Thema", "gist": "Was hier steht.", "start": "",
         "refs": [f"SZ{n}" for n in numbers[:cut]]},
        {"title": "Zweites Thema", "gist": "Was dort steht.", "start": "",
         "refs": [f"SZ{n}" for n in numbers[cut:]] or [f"SZ{numbers[-1]}"]},
    ]}


def _translate(user: str) -> dict:
    payload = json.loads(user.split("\n\n", 1)[1])
    return {k: "EN " + v["text"] for k, v in payload.items()}


def _frame() -> dict:
    return {"lead": LEAD, "gist": list(GIST), "formula": {"key": "y = a + bx", "text": "Die Gerade."},
            "design": "chapters", "hero_prompt": "a calm desk", "hero_alt": "Ein Schreibtisch."}


def _prose(user: str) -> str:
    titles = re.findall(r"^HEADING: (.+)$", user, re.M) or re.findall(r"^YOUR HEADING: (.+)$", user, re.M)
    body = ("Dieser Abschnitt erklärt seinen Gegenstand. Er nennt die Gründe. "
            "Er bleibt beim Material.")
    if len(titles) <= 1:
        return body
    return "\n\n".join(f"## {t}\n{body}" for t in titles)


def _png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (10, 70, 100)).save(buffer, "PNG")
    return buffer.getvalue()


class FakeServer:
    """Answers every call of the pipeline by the system prompt it was sent."""

    def __init__(self):
        self.requests = []
        self.images = []
        self.fail_images = False
        self.status = 200
        # Seconds every answer waits, so a test can watch a run while it is going.
        self.delay = 0.0
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def _send(self, payload: dict, status: int = 200) -> None:
                data = json.dumps(payload).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def do_POST(self):
                if outer.delay:
                    time.sleep(outer.delay)
                length = int(self.headers.get("Content-Length", 0))
                body = json.loads(self.rfile.read(length))
                outer.requests.append({"path": self.path, "body": body,
                                       "headers": {k.lower(): v for k, v in self.headers.items()}})
                if self.path.endswith("/images/generations"):
                    outer.images.append(body)
                    if outer.fail_images:
                        return self._send({"error": {"message": "no picture"}}, 500)
                    return self._send({"data": [{"b64_json": base64.b64encode(_png()).decode()}]})
                if outer.status != 200:
                    return self._send({"error": {"message": "broken"}}, outer.status)
                system = body["messages"][0]["content"]
                user = body["messages"][-1]["content"]
                if "You read a teacher's instruction" in system:
                    content = json.dumps({"way": "manner"})
                elif "You map a text" in system:
                    content = json.dumps(_topics(user))
                elif "You plan a document out of numbered material" in system:
                    content = json.dumps(_outline(user))
                elif "You translate the texts" in system:
                    content = json.dumps(_translate(user))
                elif "You write the frame" in system:
                    content = json.dumps(_frame())
                elif "You describe the one picture" in system:
                    content = "a calm desk"
                elif "You turn passages of teaching material into" in system:
                    content = json.dumps(_blocks(user))
                else:
                    content = _prose(user)
                self._send({
                    "id": "x", "object": "chat.completion", "created": 0, "model": body["model"],
                    "choices": [{"index": 0, "finish_reason": "stop",
                                 "message": {"role": "assistant", "content": content}}],
                })

        self.httpd = HTTPServer(("127.0.0.1", 0), Handler)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    @property
    def base_url(self) -> str:
        return f"http://127.0.0.1:{self.httpd.server_port}/v1"

    def names(self) -> list[str]:
        return [r["path"] for r in self.requests]

    def stop(self) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()


@pytest.fixture
def server():
    s = FakeServer()
    yield s
    s.stop()


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for name in list(os.environ):
        if name.startswith("DOC2PAGES_") or name == "OPENAI_API_KEY":
            monkeypatch.delenv(name)


@pytest.fixture
def material(tmp_path):
    """A text file long enough for a plan with several sections."""
    text = "\n\n".join(
        f"Abschnitt {i}. Die Gerade entsteht aus den Messwerten und nicht durch Probieren. "
        f"Der Schwerpunkt der Punktwolke trägt sie, weil sich die Abweichungen dort ausgleichen. "
        f"Die Steigung beträgt {i + 1} und der Achsenabschnitt {i}."
        for i in range(40))
    path = tmp_path / "regression.txt"
    path.write_text(text, encoding="utf-8")
    return path


# --- Real documents, so the readers and the layout run in the suite ---------------------

def make_pdf(pages: list[list[tuple[int, int, str]]], width: int = 595,
             height: int = 842) -> bytes:
    """A PDF with a text layer, every line placed where it is asked for."""
    font = "/Type/Font/Subtype/Type1/BaseFont/Helvetica/Encoding/WinAnsiEncoding"
    objects: list = ["<</Type/Catalog/Pages 2 0 R>>", "", f"<<{font}>>"]
    kids = []
    for i, lines in enumerate(pages):
        number = 4 + 2 * i
        kids.append(f"{number} 0 R")
        drawn = []
        for x, y, text in lines:
            safe = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            drawn.append(f"BT /F1 11 Tf 1 0 0 1 {x} {y} Tm ({safe}) Tj ET")
        stream = "\n".join(drawn).encode("cp1252")
        objects.append(f"<</Type/Page/Parent 2 0 R/MediaBox[0 0 {width} {height}]"
                       f"/Resources<</Font<</F1 3 0 R>>>>/Contents {number + 1} 0 R>>")
        objects.append(("stream", f"<</Length {len(stream)}>>".encode(), stream))
    objects[1] = f"<</Type/Pages/Kids[{' '.join(kids)}]/Count {len(pages)}>>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode()
        if isinstance(body, tuple):
            out += body[1] + b"\nstream\n" + body[2] + b"\nendstream"
        else:
            out += body.encode("cp1252")
        out += b"\nendobj\n"
    start = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<</Size {len(objects) + 1}/Root 1 0 R>>\nstartxref\n{start}\n%%EOF\n".encode()
    return bytes(out)


def _lecture_pages(count: int = 5) -> list[list[tuple[int, int, str]]]:
    """A lecture handout: a running header with a page counter and a footer on every page."""
    pages = []
    for n in range(1, count + 1):
        pages.append([
            (72, 800, f"Vorlesung Statistik | Seite {n} von {count}"),
            (72, 740, f"Kapitel {n}: Die Messwerte"),
            (72, 710, f"Abschnitt {n} erklaert die Steigung {n} und den Achsenabschnitt."),
            (72, 680, f"Die Gerade {n} entsteht aus den Messwerten, nicht durch Probieren."),
            (72, 60, "Hochschule fuer Beispiele"),
        ])
    return pages


@pytest.fixture
def lecture_pdf(tmp_path):
    """A five page PDF with a running header, a page counter and a footer."""
    path = tmp_path / "vorlesung.pdf"
    path.write_bytes(make_pdf(_lecture_pages()))
    return path


@pytest.fixture
def lecture_docx(tmp_path):
    from docx import Document

    path = tmp_path / "vorlesung.docx"
    doc = Document()
    doc.add_heading("Lineare Regression", level=1)
    for n in range(1, 4):
        doc.add_paragraph(f"Abschnitt {n} erklaert die Steigung und den Achsenabschnitt.")
    doc.save(str(path))
    return path


@pytest.fixture
def lecture_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "vorlesung.pptx"
    deck = Presentation()
    blank = deck.slide_layouts[6]
    for n in range(1, 6):
        slide = deck.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
        frame = box.text_frame
        frame.text = f"Folie {n}: Die Messwerte"
        frame.add_paragraph().text = f"Die Steigung betraegt {n} und der Achsenabschnitt {n - 1}."
        # The same line on every slide, that is what the boilerplate check is for.
        frame.add_paragraph().text = "Hochschule fuer Beispiele"
    deck.save(str(path))
    return path


@pytest.fixture
def lecture_srt(tmp_path):
    path = tmp_path / "vorlesung.srt"
    blocks = []
    for n in range(1, 6):
        start, end = 4 * (n - 1), 4 * n
        blocks.append(
            f"{n}\n00:00:{start:02d},000 --> 00:00:{end:02d},000\n"
            f"Die Steigung betraegt {n} und der Achsenabschnitt {n - 1}.\n")
    path.write_text("\n".join(blocks), encoding="utf-8")
    return path


@pytest.fixture
def lecture_vtt(tmp_path):
    path = tmp_path / "vorlesung.vtt"
    blocks = ["WEBVTT", "", "NOTE a note that says nothing about the content", ""]
    for n in range(1, 6):
        start, end = 4 * (n - 1), 4 * n
        blocks.append(f"00:00:{start:02d}.000 --> 00:00:{end:02d}.000")
        blocks.append(f"Die Gerade {n} entsteht aus den Messwerten.")
        blocks.append("")
    path.write_text("\n".join(blocks), encoding="utf-8")
    return path
